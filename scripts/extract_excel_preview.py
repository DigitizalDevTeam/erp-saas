import sys
from pathlib import Path

import openpyxl


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: python scripts/extract_excel_preview.py <path-to-xlsx/xlsm> [max_rows] [max_cols]")
        return 2

    path = Path(sys.argv[1])
    max_rows = int(sys.argv[2]) if len(sys.argv) >= 3 else 60
    max_cols = int(sys.argv[3]) if len(sys.argv) >= 4 else 16

    wb = openpyxl.load_workbook(path, data_only=False, keep_vba=str(path).lower().endswith(".xlsm"))
    print("sheets:", wb.sheetnames)

    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\n== {name} ==")
        print("dims:", ws.dimensions)

        for r in range(1, max_rows + 1):
            row = []
            for c in range(1, max_cols + 1):
                v = ws.cell(r, c).value
                if v is None:
                    v = ""
                row.append(str(v)[:80])
            if any(x != "" for x in row):
                print(f"R{r}:", row)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())

